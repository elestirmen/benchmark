#!/usr/bin/env python
"""Pipeline genelinde TEK Excel raporu: tum gorevler + RAW baseline bir arada.

Her benchmark gorevi kendi outputs/<gorev>/benchmark_results.xlsx dosyasini
uretir. Bu modul onlarin ustune, o ana kadar tamamlanmis HER SEYI tek dosyada
gosteren TUM_SONUCLAR.xlsx'i yazar.

Neden ayri bir dosya: gorev basina Excel, 15 isli bir sweepte 15 ayri dosya
demektir ve modeller arasi karsilastirma icin hepsini tek tek acmak gerekir.
Bu rapor her gorev bitiminde yenilenir, yani kosu yarim kalsa bile acilip o
ana kadarki tablo gorulebilir.

RAW baseline pipeline icinde bir kez hesaplanir ve ortak kullanilir
(end_to_end_pipeline.refresh_combined_summary); bu raporda her zaman
"RAW BASELINE" gorev adiyla gorunur, yani karsilastirma barajini gormek icin
ayri bir dosya acmak gerekmez.

Sayfalar:
    Ozet          global arama, basari sirali, RAW dahil  <- once buraya bakin
    Tum Sonuclar  her yon x senaryo x arama modu x model
    Marj Kapisi   benchmark oncesi CPU vekili (false_peak_ratio)
    Gorev Durumu  hangi is tamamlandi, hangi konfigurasyonla
"""

from __future__ import annotations

import json
import math
from pathlib import Path
from typing import Any, Iterable, Sequence

from openpyxl import Workbook

from build_benchmark_excel_openpyxl import (
    SUMMARY_COLUMNS,
    SUMMARY_FORMATS,
    display_direction,
    display_query_variant,
    display_search_mode,
    save_workbook_safely,
    set_widths,
    write_table,
)


COMBINED_WORKBOOK_NAME = "TUM_SONUCLAR.xlsx"
RAW_MODEL_ID = "RAW_BASELINE"
RAW_JOB_LABEL = "RAW BASELINE"

# Ozet sayfasi ana benchmark ile ayni konvansiyonu izler: yalniz global arama.
# ROI sonuclari Tum Sonuclar sayfasinda ayri tutulur (benchmark/README.md).
OVERVIEW_SEARCH_MODE = "global"

OVERVIEW_COLUMNS = [
    "pipeline_job",
    "model_label",
    "direction",
    "query_variant",
    "total_queries",
    "success_25m",
    "auc_25m",
    "median_error_under_25m",
    "p90_error_m",
    "coverage",
    "margin_gate_ratio",
    "margin_gate_bar",
]

OVERVIEW_HEADERS = {
    "pipeline_job": "Görev",
    "model_label": "Model / checkpoint",
    "direction": "Yön",
    "query_variant": "Senaryo",
    "total_queries": "N",
    "success_25m": "Başarı ≤25 m",
    "auc_25m": "AUC@25 m",
    "median_error_under_25m": "Medyan hata ≤25 m (m)",
    "p90_error_m": "P90 hata (m)",
    "coverage": "Kapsam",
    "margin_gate_ratio": "Marj oranı",
    "margin_gate_bar": "Marj barajı (RAW)",
}

OVERVIEW_FORMATS = {
    "success_25m": "0.0%",
    "auc_25m": "0.0%",
    "coverage": "0.0%",
    "median_error_under_25m": "0.00",
    "p90_error_m": "0.00",
    "margin_gate_ratio": "0.00",
    "margin_gate_bar": "0.00",
}

GATE_COLUMNS = [
    "pipeline_job",
    "checkpoint",
    "false_peak_ratio",
    "raw_false_peak_ratio",
    "verdict",
    "separation_z",
    "truth_ncc_mean",
    "wrong_ncc_mean",
    "margin_mean",
    "params",
    "aoi",
    "samples",
]

GATE_HEADERS = {
    "pipeline_job": "Görev",
    "checkpoint": "Checkpoint",
    "false_peak_ratio": "Oran (küçük iyi)",
    "raw_false_peak_ratio": "RAW barajı",
    "verdict": "Sonuç",
    "separation_z": "z",
    "truth_ncc_mean": "Gerçek NCC",
    "wrong_ncc_mean": "Yanlış NCC",
    "margin_mean": "Marj",
    "params": "Parametre",
    "aoi": "Bölge",
    "samples": "Örnek",
}

GATE_FORMATS = {
    "false_peak_ratio": "0.00",
    "raw_false_peak_ratio": "0.00",
    "separation_z": "0.00",
    "truth_ncc_mean": "+0.0000",
    "wrong_ncc_mean": "+0.0000",
    "margin_mean": "+0.0000",
    "params": "#,##0",
}

STATUS_COLUMNS = [
    "position",
    "name",
    "state",
    "MODEL_TYPE",
    "FILTER_COUNT",
    "KERNEL_SIZE",
    "FLAT_DEPTH",
    "HIDDEN_ACTIVATION",
    "OUTPUT_ACTIVATION",
    "LOSS_FUNCTION",
    "EPOCHS",
    "estimated_hours",
]

STATUS_HEADERS = {
    "position": "#",
    "name": "Görev",
    "state": "Durum",
    "MODEL_TYPE": "Mimari",
    "FILTER_COUNT": "f",
    "KERNEL_SIZE": "k",
    "FLAT_DEPTH": "d",
    "HIDDEN_ACTIVATION": "Gizli akt.",
    "OUTPUT_ACTIVATION": "Çıkış akt.",
    "LOSS_FUNCTION": "Kayıp",
    "EPOCHS": "Epoch",
    "estimated_hours": "Tahmini saat",
}

STATUS_FORMATS = {"estimated_hours": "0.0"}


def finite_number(value: Any) -> float | None:
    """Excel'e yazilamayan inf/nan degerlerini bosluga cevir."""
    if value is None or isinstance(value, bool):
        return None
    try:
        number = float(value)
    except (TypeError, ValueError):
        return None
    if not math.isfinite(number):
        return None
    return number


def read_json(path: Path, default: Any = None) -> Any:
    try:
        return json.loads(path.read_text(encoding="utf-8"))
    except (OSError, ValueError):
        return default


def job_display_name(row: dict[str, Any]) -> str:
    """RAW satirlari icin gorev sutununu okunur bir etikete cevir."""
    if str(row.get("model_id") or "") == RAW_MODEL_ID:
        return RAW_JOB_LABEL
    return str(row.get("pipeline_job") or "")


def model_display_label(row: dict[str, Any]) -> str:
    if str(row.get("model_id") or "") == RAW_MODEL_ID:
        return "RAW (ham görüntü)"
    label = row.get("model_label") or row.get("model_id") or ""
    return str(label)


def load_margin_gate_rows(benchmark_root: Path) -> list[dict[str, Any]]:
    """_margin_gate/*.json dosyalarini tek tabloya ac.

    Kapi bilgilendirme amaclidir ve gorevi durdurmaz (end_to_end_pipeline.
    run_margin_gate), bu yuzden zayif bir adayin benchmark'a girmis olmasi
    normaldir; tablo bunu gorunur kilar.
    """
    gate_dir = benchmark_root / "_margin_gate"
    if not gate_dir.is_dir():
        return []
    rows: list[dict[str, Any]] = []
    for path in sorted(gate_dir.glob("*.json")):
        report = read_json(path, {})
        if not isinstance(report, dict):
            continue
        job_name = path.stem
        bar = finite_number(report.get("raw_false_peak_ratio"))
        passed = set(report.get("passed") or [])
        results = report.get("results") or {}
        if not isinstance(results, dict):
            continue
        for checkpoint, entry in results.items():
            if not isinstance(entry, dict):
                continue
            if "error" in entry:
                rows.append(
                    {
                        "pipeline_job": job_name,
                        "checkpoint": checkpoint,
                        "verdict": f"HATA: {str(entry['error']).splitlines()[0][:80]}",
                        "raw_false_peak_ratio": bar,
                        "aoi": report.get("aoi"),
                        "samples": report.get("samples"),
                    }
                )
                continue
            ratio = finite_number(entry.get("false_peak_ratio"))
            if checkpoint == "RAW":
                verdict = "baraj"
            elif checkpoint in passed:
                verdict = "GEÇTİ"
            else:
                verdict = "elendi"
            rows.append(
                {
                    "pipeline_job": job_name,
                    "checkpoint": checkpoint,
                    "false_peak_ratio": ratio,
                    "raw_false_peak_ratio": bar,
                    "verdict": verdict,
                    "separation_z": finite_number(entry.get("separation_z")),
                    "truth_ncc_mean": finite_number(entry.get("truth_ncc_mean")),
                    "wrong_ncc_mean": finite_number(entry.get("wrong_ncc_mean")),
                    "margin_mean": finite_number(entry.get("margin_mean")),
                    "params": finite_number(entry.get("params")),
                    "aoi": report.get("aoi"),
                    "samples": report.get("samples"),
                }
            )
    rows.sort(
        key=lambda row: (
            row.get("false_peak_ratio") is None,
            row.get("false_peak_ratio") if row.get("false_peak_ratio") is not None else 0.0,
        )
    )
    return rows


def best_gate_ratio_by_job(gate_rows: Sequence[dict[str, Any]]) -> dict[str, dict[str, float]]:
    """Gorev basina en iyi (en kucuk) oran ve o kosunun RAW baraji."""
    best: dict[str, dict[str, float]] = {}
    for row in gate_rows:
        if row.get("checkpoint") == "RAW":
            continue
        ratio = row.get("false_peak_ratio")
        if ratio is None:
            continue
        job_name = str(row.get("pipeline_job") or "")
        current = best.get(job_name)
        if current is None or ratio < current["ratio"]:
            best[job_name] = {
                "ratio": float(ratio),
                "bar": row.get("raw_false_peak_ratio"),
            }
    return best


def build_overview_rows(
    combined_rows: Sequence[dict[str, Any]],
    gate_by_job: dict[str, dict[str, float]],
) -> list[dict[str, Any]]:
    """Global arama sonuclarini basariya gore sirala; RAW baraji da tabloda."""
    rows: list[dict[str, Any]] = []
    for source in combined_rows:
        if str(source.get("search_mode") or "") != OVERVIEW_SEARCH_MODE:
            continue
        job_name = str(source.get("pipeline_job") or "")
        gate = gate_by_job.get(job_name, {})
        rows.append(
            {
                "pipeline_job": job_display_name(source),
                "model_label": model_display_label(source),
                "direction": display_direction(source.get("direction")),
                "query_variant": display_query_variant(source.get("query_variant")),
                "total_queries": source.get("total_queries"),
                "success_25m": finite_number(source.get("success_25m")),
                "auc_25m": finite_number(source.get("auc_25m")),
                "median_error_under_25m": finite_number(
                    source.get("median_error_under_25m")
                ),
                "p90_error_m": finite_number(source.get("p90_error_m")),
                "coverage": finite_number(source.get("coverage")),
                "margin_gate_ratio": gate.get("ratio"),
                "margin_gate_bar": gate.get("bar"),
            }
        )
    rows.sort(
        key=lambda row: (
            row["success_25m"] is None,
            -(row["success_25m"] or 0.0),
            row["median_error_under_25m"] is None,
            row["median_error_under_25m"] or 0.0,
        )
    )
    return rows


def build_detail_rows(combined_rows: Sequence[dict[str, Any]]) -> list[dict[str, Any]]:
    rows: list[dict[str, Any]] = []
    for source in combined_rows:
        row = dict(source)
        row["pipeline_job"] = job_display_name(source)
        row["direction"] = display_direction(source.get("direction"))
        row["query_variant"] = display_query_variant(source.get("query_variant"))
        row["search_mode"] = display_search_mode(source.get("search_mode"))
        for key, value in list(row.items()):
            if isinstance(value, float) and not math.isfinite(value):
                row[key] = None
        rows.append(row)
    rows.sort(
        key=lambda row: (
            str(row.get("pipeline_job") or ""),
            str(row.get("direction") or ""),
            str(row.get("query_variant") or ""),
            str(row.get("search_mode") or ""),
        )
    )
    return rows


def detail_columns(rows: Sequence[dict[str, Any]]) -> list[str]:
    leading = ["pipeline_job", "search_mode"]
    columns = list(leading)
    for name in SUMMARY_COLUMNS:
        if name not in columns:
            columns.append(name)
    for row in rows:
        for name in row:
            if name not in columns:
                columns.append(name)
    return columns


def write_combined_workbook(
    benchmark_root: Path,
    combined_rows: Sequence[dict[str, Any]],
    *,
    run_id: str | None = None,
    status_rows: Sequence[dict[str, Any]] = (),
) -> tuple[Path, str]:
    """TUM_SONUCLAR.xlsx yaz; Excel acikken kilitli kopyaya duser.

    save_workbook_safely, hedef dosya Excel'de acik oldugunda zaman damgali
    bir kopyaya yazar ve kosuyu durdurmaz.
    """
    gate_rows = load_margin_gate_rows(benchmark_root)
    gate_by_job = best_gate_ratio_by_job(gate_rows)
    overview = build_overview_rows(combined_rows, gate_by_job)
    detail = build_detail_rows(combined_rows)
    columns = detail_columns(detail)

    workbook = Workbook()
    workbook.remove(workbook.active)

    sheet = workbook.create_sheet("Özet")
    write_table(
        sheet,
        overview,
        OVERVIEW_COLUMNS,
        table_name="OzetTum_01",
        number_formats=OVERVIEW_FORMATS,
        headers=[OVERVIEW_HEADERS[name] for name in OVERVIEW_COLUMNS],
    )
    set_widths(
        sheet,
        OVERVIEW_COLUMNS,
        overrides={"pipeline_job": 26.0, "model_label": 34.0},
    )

    sheet = workbook.create_sheet("Tüm Sonuçlar")
    write_table(
        sheet,
        detail,
        columns,
        table_name="TumSonuclar_01",
        number_formats=SUMMARY_FORMATS,
    )
    set_widths(sheet, columns, overrides={"pipeline_job": 26.0})

    sheet = workbook.create_sheet("Marj Kapısı")
    write_table(
        sheet,
        gate_rows,
        GATE_COLUMNS,
        table_name="MarjKapisi_01",
        number_formats=GATE_FORMATS,
        headers=[GATE_HEADERS[name] for name in GATE_COLUMNS],
    )
    set_widths(
        sheet,
        GATE_COLUMNS,
        overrides={"pipeline_job": 26.0, "checkpoint": 44.0, "verdict": 16.0},
    )

    sheet = workbook.create_sheet("Görev Durumu")
    write_table(
        sheet,
        list(status_rows),
        STATUS_COLUMNS,
        table_name="GorevDurumu_01",
        number_formats=STATUS_FORMATS,
        headers=[STATUS_HEADERS[name] for name in STATUS_COLUMNS],
    )
    set_widths(sheet, STATUS_COLUMNS, overrides={"name": 30.0, "state": 18.0})

    output_path = benchmark_root / COMBINED_WORKBOOK_NAME
    actual_path, mode = save_workbook_safely(workbook, output_path)
    return actual_path, mode
