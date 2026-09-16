"""Pipeline genelindeki tek Excel raporunun (TUM_SONUCLAR.xlsx) testleri.

Rapor 170 saatlik bir sweep boyunca her gorev sonunda yenilenir; sessizce
bozulmasi, kosunun okunabilir tek ciktisini kaybettirir.
"""

import importlib.util
import json
import sys
import tempfile
import unittest
from pathlib import Path

from openpyxl import load_workbook


BENCHMARK_DIR = Path(__file__).resolve().parents[1]
PIPELINE_PATH = BENCHMARK_DIR / "end_to_end_pipeline.py"

if str(BENCHMARK_DIR) not in sys.path:
    sys.path.insert(0, str(BENCHMARK_DIR))

import combined_report  # noqa: E402


def load_pipeline_module():
    spec = importlib.util.spec_from_file_location(
        "end_to_end_pipeline_combined_tested", PIPELINE_PATH
    )
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    return module


def summary_row(model_id, direction, variant, mode, success, **extra):
    row = {
        "model_id": model_id,
        "model_label": f"label_{model_id}",
        "direction": direction,
        "query_variant": variant,
        "search_mode": mode,
        "total_queries": 1000,
        "coverage": 0.99,
        "success_25m": success,
        "auc_25m": success * 0.8,
        "median_error_under_25m": 11.0,
        "p90_error_m": 130.0,
    }
    row.update(extra)
    return row


class CombinedReportTests(unittest.TestCase):
    def test_overview_keeps_only_global_and_sorts_by_success(self):
        rows = [
            summary_row("m1", "G__TO__B", "clean", "roi_500m", 0.90, pipeline_job="a"),
            summary_row("m1", "G__TO__B", "clean", "global", 0.20, pipeline_job="a"),
            summary_row("m2", "G__TO__B", "clean", "global", 0.70, pipeline_job="b"),
        ]
        overview = combined_report.build_overview_rows(rows, {})
        self.assertEqual(2, len(overview), "ROI satiri ozete girmemeli")
        self.assertEqual([0.70, 0.20], [row["success_25m"] for row in overview])

    def test_raw_baseline_is_labelled_and_kept_in_the_overview(self):
        """RAW baraji ayni tabloda gorunmeli; ayri dosya acmak gerekmemeli."""
        rows = [
            summary_row(
                "RAW_BASELINE",
                "G__TO__B",
                "clean",
                "global",
                0.61,
                pipeline_job="shared_raw_baseline",
            ),
            summary_row("m1", "G__TO__B", "clean", "global", 0.70, pipeline_job="a"),
        ]
        overview = combined_report.build_overview_rows(rows, {})
        labels = [row["pipeline_job"] for row in overview]
        self.assertIn(combined_report.RAW_JOB_LABEL, labels)
        self.assertEqual("a", overview[0]["pipeline_job"])

    def test_non_finite_numbers_never_reach_excel(self):
        """inf/nan yazilan bir hucre dosyayi Excel'de acilamaz hale getirir."""
        rows = [
            summary_row(
                "m1",
                "G__TO__B",
                "clean",
                "global",
                0.5,
                pipeline_job="a",
                p90_error_m=float("inf"),
                median_error_under_25m=float("nan"),
            )
        ]
        overview = combined_report.build_overview_rows(rows, {})
        self.assertIsNone(overview[0]["p90_error_m"])
        self.assertIsNone(overview[0]["median_error_under_25m"])
        detail = combined_report.build_detail_rows(rows)
        self.assertIsNone(detail[0]["p90_error_m"])

    def test_margin_gate_table_reports_verdicts_and_survives_errors(self):
        with tempfile.TemporaryDirectory() as temp_dir:
            root = Path(temp_dir)
            gate_dir = root / "_margin_gate"
            gate_dir.mkdir()
            (gate_dir / "job_a.json").write_text(
                json.dumps(
                    {
                        "aoi": "all",
                        "samples": 60,
                        "raw_false_peak_ratio": 0.62,
                        "results": {
                            "RAW": {"false_peak_ratio": 0.62},
                            "good.h5": {"false_peak_ratio": 0.34},
                            "weak.h5": {"false_peak_ratio": 0.81},
                            "broken.h5": {"error": "yuklenemedi"},
                        },
                        "passed": ["good.h5"],
                        "failed": ["weak.h5", "broken.h5"],
                    }
                ),
                encoding="utf-8",
            )
            gate_rows = combined_report.load_margin_gate_rows(root)
            by_checkpoint = {row["checkpoint"]: row for row in gate_rows}
            self.assertEqual("GEÇTİ", by_checkpoint["good.h5"]["verdict"])
            self.assertEqual("elendi", by_checkpoint["weak.h5"]["verdict"])
            self.assertEqual("baraj", by_checkpoint["RAW"]["verdict"])
            self.assertTrue(by_checkpoint["broken.h5"]["verdict"].startswith("HATA"))

            # RAW satiri gorevin kendi skoru degildir; en iyi orana karismamali.
            best = combined_report.best_gate_ratio_by_job(gate_rows)
            self.assertAlmostEqual(0.34, best["job_a"]["ratio"])

    def test_workbook_has_every_sheet_and_opens(self):
        with tempfile.TemporaryDirectory() as temp_dir:
            root = Path(temp_dir)
            rows = [
                summary_row(
                    "RAW_BASELINE",
                    "G__TO__B",
                    "clean",
                    "global",
                    0.61,
                    pipeline_job="shared_raw_baseline",
                ),
                summary_row("m1", "G__TO__B", "clean", "global", 0.70, pipeline_job="a"),
            ]
            path, mode = combined_report.write_combined_workbook(
                root, rows, run_id="20260905_120000", status_rows=[]
            )
            self.assertEqual("primary", mode)
            workbook = load_workbook(path)
            self.assertEqual(
                ["Özet", "Tüm Sonuçlar", "Marj Kapısı", "Görev Durumu"],
                workbook.sheetnames,
            )

    def test_empty_input_still_produces_an_openable_workbook(self):
        """Hicbir gorev bitmemisken bile dosya acilabilir olmali.

        write_table bos veri setine bilerek bir bos satir ekler: Excel masaustu
        yalniz baslik iceren bir tabloyu reddeder. Beklenen sey bos veri degil,
        bozuk olmayan bir dosyadir.
        """
        with tempfile.TemporaryDirectory() as temp_dir:
            root = Path(temp_dir)
            path, _ = combined_report.write_combined_workbook(
                root, [], run_id="20260905_120000", status_rows=[]
            )
            self.assertTrue(path.is_file())
            workbook = load_workbook(path)
            sheet = workbook["Özet"]
            self.assertEqual(2, sheet.max_row)
            self.assertTrue(
                all(cell.value is None for cell in sheet[2]),
                "veri satiri gercekten bos olmali",
            )


class CombinedReportPipelineIntegrationTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.pipeline = load_pipeline_module()

    def write_complete_benchmark(self, output_dir, rows):
        output_dir.mkdir(parents=True, exist_ok=True)
        (output_dir / "results.jsonl").write_text("{}\n", encoding="utf-8")
        (output_dir / "summary.json").write_text(
            json.dumps(rows, ensure_ascii=False), encoding="utf-8"
        )
        (output_dir / "summary_metadata.json").write_text(
            json.dumps(
                {
                    "expected_queries_per_group": self.pipeline.BENCHMARK_MAX_QUERIES,
                    "incomplete_group_count": 0,
                    "group_count": len(rows),
                }
            ),
            encoding="utf-8",
        )
        (output_dir / "summary_incomplete.json").write_text("[]", encoding="utf-8")
        (output_dir / "benchmark_results.xlsx").write_text("x", encoding="utf-8")
        (output_dir / "excel_validation.json").write_text(
            json.dumps({"zip_integrity": "ok", "formula_error_literals": []}),
            encoding="utf-8",
        )

    def test_refresh_writes_the_workbook_without_a_raw_baseline(self):
        """Ilk gorev bitmeden RAW yoktur; rapor yine de acilabilmeli."""
        with tempfile.TemporaryDirectory() as temp_dir:
            root = Path(temp_dir)
            self.pipeline.PIPELINE_RUN_ID = "20260905_120000"
            self.pipeline.PIPELINE_MODELS_ROOT = root / "models"
            self.pipeline.PIPELINE_BENCHMARK_ROOT = root / "outputs"
            self.pipeline.PIPELINE_MODELS_ROOT.mkdir(parents=True)
            self.pipeline.PIPELINE_BENCHMARK_ROOT.mkdir(parents=True)

            first = self.pipeline.TRAINING_JOBS[0]["name"]
            self.write_complete_benchmark(
                self.pipeline.PIPELINE_BENCHMARK_ROOT / first,
                [summary_row("m1", "G__TO__B", "clean", "global", 0.70)],
            )

            json_path = self.pipeline.refresh_combined_summary(None, {first})
            self.assertIsNotNone(json_path)
            payload = json.loads(json_path.read_text(encoding="utf-8"))
            self.assertIsNone(payload["shared_raw_source"])
            self.assertEqual(0, payload["raw_baseline_row_count"])
            self.assertEqual(1, payload["model_row_count"])
            self.assertTrue(
                (self.pipeline.PIPELINE_BENCHMARK_ROOT / "TUM_SONUCLAR.xlsx").is_file()
            )

    def test_job_status_rows_cover_every_job_and_mark_selection(self):
        with tempfile.TemporaryDirectory() as temp_dir:
            root = Path(temp_dir)
            self.pipeline.PIPELINE_MODELS_ROOT = root / "models"
            self.pipeline.PIPELINE_BENCHMARK_ROOT = root / "outputs"
            self.pipeline.PIPELINE_MODELS_ROOT.mkdir(parents=True)
            self.pipeline.PIPELINE_BENCHMARK_ROOT.mkdir(parents=True)

            first = self.pipeline.TRAINING_JOBS[0]["name"]
            rows = self.pipeline.build_job_status_rows({first})
            self.assertEqual(len(self.pipeline.TRAINING_JOBS), len(rows))
            self.assertEqual("baslamadi", rows[0]["state"])
            self.assertIn("secilmedi", rows[1]["state"])

    def test_report_failure_never_stops_the_run(self):
        """11 saatlik bir egitimin ardindan rapor yazimi yuzunden cikilmamali."""
        self.pipeline.PIPELINE_BENCHMARK_ROOT = Path("Z:/olmayan/yol/xyz")
        result = self.pipeline.write_combined_excel(
            [summary_row("m1", "G__TO__B", "clean", "global", 0.5)], None
        )
        self.assertIsNone(result)


if __name__ == "__main__":
    unittest.main()
